from __future__ import annotations

import tempfile
import sys
import threading
import time
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from dispatcher.history import HistoryStore
from dispatcher.models import ScheduleConfig, TaskDefinition
from dispatcher.modules.broadcast import InternalBotClient, export_marker_blocks, group_news, parse_cue, run_news, run_universal_cut
from dispatcher.modules.common import TaskContext
from dispatcher.modules.registry import defaults_for
from dispatcher.modules.music import run_foreign_agents, run_list_mp3
from dispatcher.modules.reports import run_excel_report
from dispatcher.modules.web_tools import run_make_links
from dispatcher.schedule_engine import next_run, occurrence_key, schedule_summary
from dispatcher.task_runner import TaskRunner


def context(task: TaskDefinition, root: Path) -> TaskContext:
    return TaskContext(
        task=task,
        settings={"timezone": "Europe/Moscow"},
        secrets={},
        run_dir=root / "run",
        cancel_event=threading.Event(),
        log=lambda _message: None,
        progress=lambda _value, _message: None,
    )


class ScheduleTests(unittest.TestCase):
    def test_multiple_times_and_next_run(self):
        schedule = ScheduleConfig(
            mode="times",
            weekdays=[0, 1, 2, 3, 4],
            times=["08:05", "12:30"],
        )
        now = datetime(2026, 7, 27, 8, 4)  # Monday
        self.assertEqual(next_run(schedule, now), datetime(2026, 7, 27, 8, 5))
        self.assertEqual(
            occurrence_key(schedule, datetime(2026, 7, 27, 12, 30)),
            "2026-07-27T12:30",
        )
        self.assertIn("Пн–Пт", schedule_summary(schedule))

    def test_interval(self):
        schedule = ScheduleConfig(
            mode="interval",
            weekdays=list(range(7)),
            interval_minutes=30,
            window_start="08:00",
            window_end="09:00",
        )
        self.assertIsNotNone(occurrence_key(schedule, datetime(2026, 7, 26, 8, 30)))
        self.assertIsNone(occurrence_key(schedule, datetime(2026, 7, 26, 8, 45)))


class BroadcastDefaultsTests(unittest.TestCase):
    def test_news_bot_values_are_separate_per_station(self):
        news = defaults_for("broadcast.news")

        self.assertEqual(news["internal_bot_message_kometa"], "НОВОСТИ КОМЕТА {date}")
        self.assertEqual(news["internal_bot_filename_kometa"], "kometa")
        self.assertEqual(news["internal_bot_message_stolica"], "НОВОСТИ СТОЛИЦА {date}")
        self.assertEqual(news["internal_bot_filename_stolica"], "stolica")
        self.assertEqual(news["internal_bot_message_briz"], "НОВОСТИ БРИЗ {date}")
        self.assertEqual(news["internal_bot_filename_briz"], "briz")
        self.assertEqual(news["tormozok_name"], "Тормозок")
        self.assertEqual(news["interview_name"], "Интервью")
        self.assertEqual(news["internal_bot_filename_tormozok"], "tormozok")
        self.assertEqual(news["internal_bot_filename_interview"], "interview")

    def test_program_fragment_bot_defaults_stay_compatible(self):
        shift = defaults_for("broadcast.morning_shift")
        guest = defaults_for("broadcast.morning_guest")
        self.assertEqual(shift["internal_bot_message"], "УТРЕННЯЯ СМЕНА {date}")
        self.assertEqual(shift["internal_bot_filename"], "{type}_{source}{ext}")
        self.assertEqual(guest["internal_bot_message"], "УТРЕННИЙ ГОСТЬ {date}")
        self.assertEqual(guest["internal_bot_filename"], "{type}_{source}{ext}")


class UniversalCutTests(unittest.TestCase):
    def test_defaults_are_empty_and_independent(self):
        first = defaults_for("broadcast.universal_cut")
        second = defaults_for("broadcast.universal_cut")
        self.assertEqual(first["sources"], [])
        first["sources"].append({"name": "A"})
        self.assertEqual(second["sources"], [])

    def test_each_source_has_own_hours_and_filename(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            a = root / "a"
            b = root / "b"
            a.mkdir()
            b.mkdir()
            for folder, hour in ((a, "09"), (b, "12")):
                (folder / f"17-08-2026_{hour}-00.mp3").write_bytes(b"mp3")
                (folder / f"17-08-2026_{hour}-00.cue").write_text("cue", encoding="utf-8")

            calls = []

            def fake_export(ctx, mp3, cue, output, output_prefix, start_marker, end_marker=None, same_marker_twice=False):
                calls.append((mp3.parent.name, output_prefix, start_marker, end_marker, same_marker_twice))
                path = output / f"{output_prefix}_{cue.stem}.mp3"
                path.write_bytes(b"exported")
                return path

            task = TaskDefinition(
                id="universal",
                name="Универсальная",
                module="broadcast.universal_cut",
                params={
                    "date_mode": "manual",
                    "manual_date": "17-08-2026",
                    "cut_mode": "between",
                    "start_marker": "влет",
                    "end_marker": "вылет",
                    "sources": [
                        {"enabled": True, "name": "A", "folder": str(a), "hours": "9", "message": "A", "filename": "alpha"},
                        {"enabled": True, "name": "B", "folder": str(b), "hours": "12", "message": "B", "filename": "beta"},
                    ],
                    "output_dir": str(root / "out"),
                    "internal_bot_enabled": False,
                },
            )
            with patch("dispatcher.modules.broadcast.export_between_markers", side_effect=fake_export):
                result = run_universal_cut(context(task, root))

            self.assertEqual(len(result["files"]), 2)
            self.assertEqual(
                calls,
                [
                    ("a", "alpha", "влет", "вылет", False),
                    ("b", "beta", "влет", "вылет", False),
                ],
            )

    def test_single_marker_mode_uses_marker_block_export(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source"
            source.mkdir()
            (source / "17-08-2026_10-00.mp3").write_bytes(b"mp3")
            (source / "17-08-2026_10-00.cue").write_text("cue", encoding="utf-8")
            calls = []

            def fake_export(ctx, mp3, cue, output, output_prefix, marker):
                calls.append((output_prefix, marker))
                path = output / f"{output_prefix}_{cue.stem}.mp3"
                path.write_bytes(b"exported")
                return [path]

            task = TaskDefinition(
                id="track",
                name="Track",
                module="broadcast.universal_cut",
                params={
                    "date_mode": "manual",
                    "manual_date": "17-08-2026",
                    "cut_mode": "track",
                    "same_marker": "новости",
                    "sources": [{"name": "Test", "folder": str(source), "hours": "10", "filename": "test"}],
                    "output_dir": str(root / "out"),
                    "internal_bot_enabled": False,
                },
            )
            with patch("dispatcher.modules.broadcast.export_marker_blocks", side_effect=fake_export):
                run_universal_cut(context(task, root))
            self.assertEqual(calls, [("test", "новости")])

    def test_legacy_same_mode_is_migrated_to_marker_block_logic(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source"
            source.mkdir()
            (source / "17-08-2026_10-00.mp3").write_bytes(b"mp3")
            (source / "17-08-2026_10-00.cue").write_text("cue", encoding="utf-8")
            with patch("dispatcher.modules.broadcast.export_marker_blocks", return_value=[]) as fake_export:
                task = TaskDefinition(
                    id="legacy",
                    name="Legacy",
                    module="broadcast.universal_cut",
                    params={
                        "date_mode": "manual",
                        "manual_date": "17-08-2026",
                        "cut_mode": "same",
                        "same_marker": "новости",
                        "sources": [{"name": "Test", "folder": str(source), "hours": "10", "filename": "test"}],
                        "output_dir": str(root / "out"),
                        "internal_bot_enabled": False,
                    },
                )
                run_universal_cut(context(task, root))
            fake_export.assert_called_once()

    def test_legacy_combined_mode_falls_back_to_between(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source"
            source.mkdir()
            (source / "17-08-2026_10-00.mp3").write_bytes(b"mp3")
            (source / "17-08-2026_10-00.cue").write_text("cue", encoding="utf-8")
            calls = []

            def fake_export(ctx, mp3, cue, output, prefix, start, end, same_marker_twice=False):
                calls.append((prefix, start, end))
                path = output / f"{prefix}_{cue.stem}.mp3"
                path.write_bytes(b"exported")
                return path

            task = TaskDefinition(
                id="legacy-combined", name="Legacy combined", module="broadcast.universal_cut",
                params={
                    "date_mode": "manual", "manual_date": "17-08-2026",
                    "cut_mode": "combined",
                    "start_marker": "влет", "end_marker": "вылет", "same_marker": "новости",
                    "sources": [{"name": "Test", "folder": str(source), "hours": "10", "filename": "test"}],
                    "output_dir": str(root / "out"), "internal_bot_enabled": False,
                },
            )
            with patch("dispatcher.modules.broadcast.export_between_markers", side_effect=fake_export):
                run_universal_cut(context(task, root))
            self.assertEqual(calls, [("test", "влет", "вылет")])

    def test_single_marker_groups_consecutive_matching_tracks(self):
        class FakeSlice:
            def __init__(self, exported):
                self.exported = exported

            def export(self, destination, format):
                self.exported.append(Path(destination).name)

        class FakeAudio:
            exported = []

            @classmethod
            def from_mp3(cls, _path):
                cls.exported = []
                return cls()

            def __len__(self):
                return 60 * 60 * 1000

            def __getitem__(self, _slice):
                return FakeSlice(type(self).exported)

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            mp3 = root / "17-08-2026_11-00.mp3"
            cue = root / "17-08-2026_11-00.cue"
            mp3.write_bytes(b"mp3")
            cue.write_text("cue", encoding="utf-8")
            tracks = [
                {"title": "Новости", "start": 0.0, "end": 30.0},
                {"title": "Новости погода", "start": 30.0, "end": 60.0},
                {"title": "Музыка", "start": 60.0, "end": 1200.0},
                {"title": "Новости", "start": 1200.0, "end": 1260.0},
                {"title": "Музыка", "start": 1260.0, "end": None},
            ]
            task = TaskDefinition(id="x", name="X", module="broadcast.universal_cut", params={})
            with patch("dispatcher.modules.broadcast._audio_segment", return_value=FakeAudio), \
                 patch("dispatcher.modules.broadcast.parse_cue", return_value=tracks):
                made = export_marker_blocks(
                    context(task, root), mp3, cue, root, "news", "новости"
                )
            self.assertEqual(
                [item.name for item in made],
                ["news_17-08-2026_11-00.mp3", "news_17-08-2026_11-20.mp3"],
            )
            self.assertEqual(FakeAudio.exported, [item.name for item in made])
            blocks = group_news(tracks, "новости")
            self.assertEqual([len(block) for block in blocks], [2, 1])


class CueTests(unittest.TestCase):
    def test_parse_and_group(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "test.cue"
            path.write_text(
                """
TRACK 01 AUDIO
TITLE "Новости начало"
INDEX 01 00:00:00
TRACK 02 AUDIO
TITLE "Новости продолжение"
INDEX 01 01:00:00
TRACK 03 AUDIO
TITLE "Музыка"
INDEX 01 02:00:00
""".strip(),
                encoding="utf-8",
            )
            tracks = parse_cue(path)
            self.assertEqual(len(tracks), 3)
            blocks = group_news(tracks, "новости")
            self.assertEqual(len(blocks), 1)
            self.assertEqual(len(blocks[0]), 2)


class InternalBotTests(unittest.TestCase):
    def test_readable_default_filename_template_reconstructs_original(self):
        from dispatcher.modules.broadcast import _internal_bot_upload_name

        for name in (
            "smena_17-08-2026_08-00.mp3",
            "guest_17-08-2026_10-00.mp3",
            "kometa_17-08-2026_11-00.mp3",
            "tormozok_17-08-2026_12-00.mp3",
        ):
            path = Path(name)
            rendered = _internal_bot_upload_name(
                "{type}_{source}{ext}",
                path,
                date_text="17-08-2026",
                task_name="test",
                group_title="test",
                index=1,
            )
            self.assertEqual(rendered, name)

    def test_human_filename_base_preserves_source_date_and_time(self):
        from dispatcher.modules.broadcast import _internal_bot_upload_name_from_base

        path = Path("stolica_17-08-2026_11-00.mp3")
        self.assertEqual(
            _internal_bot_upload_name_from_base("novosti_stolica", path),
            "novosti_stolica_17-08-2026_11-00.mp3",
        )
        self.assertEqual(
            _internal_bot_upload_name_from_base("Новости Столица.mp3", path),
            "Новости Столица_17-08-2026_11-00.mp3",
        )

    def test_internal_endpoints_and_payload(self):
        class Response:
            status_code = 200
            text = '{"ok": true}'

            @staticmethod
            def json():
                return {"ok": True}

        class Requests:
            def __init__(self):
                self.calls = []

            def post(self, url, **kwargs):
                self.calls.append((url, kwargs))
                return Response()

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "fragment.mp3"
            audio.write_bytes(b"test")
            task = TaskDefinition(id="bot", name="Bot", module="broadcast.news")
            ctx = TaskContext(
                task=task,
                settings={
                    "internal_bot_base_url": "http://192.168.2.101:8000",
                    "internal_bot_channel": "extracted_media",
                    "internal_bot_title": "Вырезки из эфира",
                    "internal_bot_max_retries": 1,
                    "internal_bot_retry_delay": 1,
                    "internal_bot_file_delay": 0,
                },
                secrets={"internal_bot_key": "secret"},
                run_dir=root,
                cancel_event=threading.Event(),
                log=lambda _message: None,
                progress=lambda _value, _message: None,
            )
            requests = Requests()
            with patch("dispatcher.modules.broadcast.need", return_value=requests):
                client = InternalBotClient(ctx)
            client.send_group("НОВОСТИ КОМЕТА 28-07-2026", [audio])

            self.assertEqual(
                [call[0] for call in requests.calls],
                [
                    "http://192.168.2.101:8000/api/internal/message",
                    "http://192.168.2.101:8000/api/internal/upload",
                ],
            )
            self.assertEqual(requests.calls[0][1]["json"]["channel"], "extracted_media")
            self.assertEqual(requests.calls[1][1]["data"]["key"], "secret")
            self.assertIn("file", requests.calls[1][1]["files"])
            self.assertEqual(requests.calls[1][1]["files"]["file"][0], "fragment.mp3")

    def test_custom_message_and_upload_filename(self):
        class Response:
            status_code = 200
            text = '{"ok": true}'

            @staticmethod
            def json():
                return {"ok": True}

        class Requests:
            def __init__(self):
                self.calls = []

            def post(self, url, **kwargs):
                self.calls.append((url, kwargs))
                return Response()

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "fragment.mp3"
            audio.write_bytes(b"test")
            task = TaskDefinition(id="bot", name="Новости 11:00", module="broadcast.news")
            ctx = TaskContext(
                task=task,
                settings={
                    "internal_bot_base_url": "http://192.168.2.101:8000",
                    "internal_bot_channel": "extracted_media",
                    "internal_bot_title": "Вырезки из эфира",
                    "internal_bot_max_retries": 1,
                    "internal_bot_retry_delay": 1,
                    "internal_bot_separator_delay": 0,
                    "internal_bot_file_delay": 0,
                },
                secrets={"internal_bot_key": "secret"},
                run_dir=root,
                cancel_event=threading.Event(),
                log=lambda _message: None,
                progress=lambda _value, _message: None,
            )
            requests = Requests()
            with patch("dispatcher.modules.broadcast.need", return_value=requests):
                client = InternalBotClient(ctx)

            client.send_group(
                "НОВОСТИ СТОЛИЦА 17-08-2026",
                [audio],
                message_template="Архив {date}: {group}",
                filename_template="news_{date}_{index}",
                date_text="17-08-2026",
            )

            self.assertEqual(
                requests.calls[0][1]["json"]["text"],
                "Архив 17-08-2026: НОВОСТИ СТОЛИЦА 17-08-2026",
            )
            self.assertEqual(
                requests.calls[1][1]["files"]["file"][0],
                "news_17-08-2026_1.mp3",
            )


class BroadcastSourceTests(unittest.TestCase):
    def test_disabled_sources_are_not_scanned(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            output = root / "cash"
            task = TaskDefinition(
                id="disabled-sources",
                name="Disabled sources",
                module="broadcast.news",
                params={
                    "use_kometa": False,
                    "use_stolica": False,
                    "use_briz": False,
                    "source_kometa": str(root / "missing-kometa"),
                    "source_stolica": str(root / "missing-stolica"),
                    "source_briz": str(root / "missing-briz"),
                    "output_dir": str(output),
                    "internal_bot_enabled": True,
                },
            )
            result = run_news(context(task, root))
            self.assertEqual(result["files"], [])
            self.assertIn("не найдено", result["message"])

    def test_manual_and_regular_hours_are_independent(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "kometa"
            source.mkdir()
            captured_hours: list[list[int]] = []

            def collect_hours(_folder, _day_text, hours):
                captured_hours.append(list(hours))
                return {}

            params = {
                "use_kometa": True,
                "use_stolica": False,
                "use_briz": False,
                "source_kometa": str(source),
                "source_stolica": str(root / "missing-stolica"),
                "source_briz": str(root / "missing-briz"),
                "hours_kometa": "9-10",
                "manual_hours_kometa": "1-2",
                "manual_date": "28-07-2026",
                "output_dir": str(root / "cash"),
                "internal_bot_enabled": False,
            }
            task = TaskDefinition(
                id="hours-profiles",
                name="Hours profiles",
                module="broadcast.news",
                params={**params, "date_mode": "manual"},
            )
            with patch(
                "dispatcher.modules.broadcast.collect_pairs",
                side_effect=collect_hours,
            ):
                run_news(context(task, root))
            self.assertEqual(captured_hours, [[1, 2]])

            for date_mode in ("today", "last_workday"):
                captured_hours.clear()
                task.params["date_mode"] = date_mode
                with patch(
                    "dispatcher.modules.broadcast.collect_pairs",
                    side_effect=collect_hours,
                ):
                    run_news(context(task, root))
                self.assertEqual(captured_hours, [[9, 10]])


class TaskRunnerTests(unittest.TestCase):
    def test_queued_task_can_be_stopped_before_worker_starts(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            task = TaskDefinition(
                id="queued",
                name="Queued",
                module="external",
                workdir=str(root),
                params={"program": str(root / "does-not-need-to-exist.py")},
            )
            runner = TaskRunner(HistoryStore(root / "history.sqlite3"))
            self.assertTrue(runner.queue(task.id))
            self.assertTrue(runner.is_active(task.id))
            self.assertTrue(runner.cancel(task.id))
            result = runner.execute(task, {}, {})
            self.assertEqual(result.status, "cancelled")
            self.assertFalse(runner.is_active(task.id))

    def test_cancel_really_terminates_running_work(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            logs = root / "logs"
            runs = root / "runs"
            logs.mkdir()
            runs.mkdir()
            script = root / "slow.py"
            marker = root / "started.txt"
            script.write_text(
                "from pathlib import Path\n"
                "import time\n"
                f"Path({str(marker)!r}).write_text('started', encoding='utf-8')\n"
                "while True:\n"
                "    time.sleep(1)\n",
                encoding="utf-8",
            )
            task = TaskDefinition(
                id="slow",
                name="Slow",
                module="external",
                workdir=str(root),
                params={
                    "program": str(script),
                    "arguments": "",
                    "python_path": sys.executable,
                    "timeout": 0,
                },
            )
            runner = TaskRunner(HistoryStore(root / "history.sqlite3"))
            result_holder = []
            with (
                patch("dispatcher.task_runner.logs_dir", return_value=logs),
                patch("dispatcher.task_runner.temp_dir", return_value=runs),
            ):
                worker = threading.Thread(
                    target=lambda: result_holder.append(runner.execute(task, {}, {})),
                    daemon=True,
                )
                worker.start()
                deadline = time.monotonic() + 5
                while not marker.exists() and time.monotonic() < deadline:
                    time.sleep(0.02)
                self.assertTrue(marker.exists())
                self.assertTrue(runner.cancel(task.id))
                worker.join(timeout=6)

            self.assertFalse(worker.is_alive())
            self.assertEqual(result_holder[0].status, "cancelled")


class FileModuleTests(unittest.TestCase):
    def test_list_and_foreign_agents(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "sub").mkdir()
            (root / "ДДТ - Тест.mp3").write_bytes(b"")
            (root / "sub" / "Other.mp3").write_bytes(b"")
            (root / "block_list.txt").write_text("ДДТ\n", encoding="utf-8")

            list_task = TaskDefinition(
                id="list",
                name="List",
                module="music.list_mp3",
                workdir=str(root),
                params={"source_dir": str(root), "output_file": "list.txt", "include_relative_path": True},
            )
            result = run_list_mp3(context(list_task, root))
            self.assertIn("2", result["message"])
            self.assertTrue((root / "list.txt").exists())

            check_task = TaskDefinition(
                id="check",
                name="Check",
                module="music.foreign_agents",
                workdir=str(root),
                params={
                    "source_dir": str(root),
                    "block_list": "block_list.txt",
                    "report_file": "matches.txt",
                },
            )
            result = run_foreign_agents(context(check_task, root))
            self.assertIn("совпадений: 1", result["message"])
            self.assertIn("ДДТ", (root / "matches.txt").read_text(encoding="utf-8"))

    def test_make_links(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_folder = root / "reports"
            source_folder.mkdir()
            (source_folder / "artists.txt").write_text("Артист - Песня\n", encoding="utf-8")
            task = TaskDefinition(
                id="links",
                name="Links",
                module="yandex.make_links",
                workdir=str(root),
                params={
                    "input_file": str(source_folder / "artists.txt"),
                    "output_file": "links.txt",
                },
            )
            run_make_links(context(task, root))
            self.assertIn(
                "music.yandex.ru/search",
                (source_folder / "links.txt").read_text(encoding="utf-8"),
            )

    def test_excel_format_only(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "report.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Название", "Общая длительность"])
            sheet.append(["Песня", "01:30"])
            book.save(source)
            task = TaskDefinition(
                id="excel",
                name="Excel",
                module="excel.report",
                workdir=str(root),
                params={
                    "input_file": str(source),
                    "mode": "format_only",
                    "output_file": str(root / "ready.xlsx"),
                    "duration_header": "Общая длительность",
                },
            )
            run_excel_report(context(task, root))
            ready = load_workbook(root / "ready.xlsx")
            self.assertEqual(ready.active["B2"].value, "00:01:30")
            self.assertEqual(ready.active["B3"].value, "00:01:30")

    def test_excel_default_output_is_next_to_input(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            reports = root / "reports"
            reports.mkdir()
            source = reports / "эфир.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Название", "Общая длительность"])
            sheet.append(["Песня", "00:30"])
            book.save(source)
            task = TaskDefinition(
                id="excel-nearby",
                name="Excel nearby",
                module="excel.report",
                params={
                    "input_file": str(source),
                    "mode": "format_only",
                    "output_file": "",
                    "duration_header": "Общая длительность",
                },
            )
            run_excel_report(context(task, root))
            self.assertTrue((reports / "эфир_готово.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
