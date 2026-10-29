from __future__ import annotations

from copy import copy
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Any

from .common import TaskContext, need, resolve_path


def _header_key(value: Any) -> str:
    if value is None:
        return ""
    return "".join(str(value).replace("\n", " ").split()).strip().lower()


def _header_map(sheet, max_column: int | None = None) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, (max_column or sheet.max_column) + 1):
        value = sheet.cell(1, column).value
        if value is not None:
            result[str(value).strip()] = column
            result[_header_key(value)] = column
    return result


def _column(headers: dict[str, int], name: str) -> int | None:
    return headers.get(name) or headers.get(_header_key(name))


def _duration_seconds(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, timedelta):
        return int(round(value.total_seconds()))
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)):
        if 0 <= float(value) < 10:
            return int(round(float(value) * 86400))
        return int(round(float(value)))
    if isinstance(value, str):
        parts = value.strip().split(":")
        if len(parts) == 2 and all(part.isdigit() for part in parts):
            return int(parts[0]) * 60 + int(parts[1])
        if len(parts) == 3 and all(part.isdigit() for part in parts):
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + int(parts[2])
    return None


def _hhmmss(seconds: int) -> str:
    seconds = max(0, int(seconds))
    return f"{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}:{seconds % 60:02d}"


def _normalize_sheet_durations(sheet) -> tuple[int, int]:
    changed = valid = 0
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if not isinstance(cell.value, str) or ":" not in cell.value:
                continue
            seconds = _duration_seconds(cell.value)
            if seconds is None:
                continue
            normalized = _hhmmss(seconds)
            if cell.value.strip() != normalized:
                cell.value = normalized
                changed += 1
            else:
                valid += 1
            cell.number_format = "@"
    return changed, valid


def _add_total(sheet, duration_column: int, max_column: int | None = None):
    from openpyxl.styles import Border, Font, PatternFill, Side

    total = 0
    counted = 0
    for row in range(2, sheet.max_row + 1):
        seconds = _duration_seconds(sheet.cell(row, duration_column).value)
        if seconds is not None:
            total += seconds
            counted += 1
    total_row = sheet.max_row + 1
    label_column = max(1, duration_column - 1)
    sheet.cell(total_row, label_column, "ИТОГО:")
    sheet.cell(total_row, duration_column, _hhmmss(total))
    thin = Side(style="thin", color="999999")
    for column in range(1, (max_column or sheet.max_column) + 1):
        cell = sheet.cell(total_row, column)
        cell.border = Border(top=thin)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.cell(total_row, label_column).font = Font(bold=True)
    sheet.cell(total_row, duration_column).font = Font(bold=True)
    sheet.cell(total_row, duration_column).number_format = "@"
    return counted, _hhmmss(total)


def _copy_cell(source, destination) -> None:
    destination.value = source.value
    if source.has_style:
        destination.font = copy(source.font)
        destination.fill = copy(source.fill)
        destination.border = copy(source.border)
        destination.alignment = copy(source.alignment)
        destination.protection = copy(source.protection)
        destination.number_format = source.number_format


def _combine_sheets(workbook):
    from openpyxl import Workbook

    if not workbook.worksheets:
        raise ValueError("В книге нет листов")
    first = workbook.worksheets[0]
    max_column = first.max_column
    while max_column > 1 and first.cell(1, max_column).value is None:
        max_column -= 1
    result = Workbook()
    output = result.active
    output.title = "Таблица"
    for column in range(1, max_column + 1):
        _copy_cell(first.cell(1, column), output.cell(1, column))
        letter = first.cell(1, column).column_letter
        output.column_dimensions[letter].width = first.column_dimensions[letter].width
    target_row = 2
    for sheet in workbook.worksheets:
        for source_row in range(2, sheet.max_row + 1):
            if all(sheet.cell(source_row, col).value is None for col in range(1, max_column + 1)):
                continue
            for column in range(1, max_column + 1):
                _copy_cell(sheet.cell(source_row, column), output.cell(target_row, column))
            target_row += 1
    return result, output, max_column


def _clean_rows(sheet, max_column: int, name_column: int, genre_column: int | None, bad_value: str):
    bad_key = _header_key(bad_value)
    removed_bad = removed_empty = removed_duplicates = 0
    seen: set[str] = set()
    keep_rows: list[int] = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, column).value for column in range(1, max_column + 1)]
        if any(_header_key(value) == bad_key for value in values):
            removed_bad += 1
            continue
        incomplete = any(
            (value is None or str(value).strip() == "") and column != genre_column
            for column, value in enumerate(values, start=1)
        )
        if incomplete:
            removed_empty += 1
            continue
        name_key = _header_key(sheet.cell(row, name_column).value)
        if name_key in seen:
            removed_duplicates += 1
            continue
        seen.add(name_key)
        keep_rows.append(row)
    for row in range(sheet.max_row, 1, -1):
        if row not in keep_rows:
            sheet.delete_rows(row, 1)
    return removed_bad, removed_empty, removed_duplicates


def run_excel_report(ctx: TaskContext):
    need("openpyxl")
    from openpyxl import load_workbook

    params = ctx.task.params
    raw_input = str(params.get("input_file", "")).strip()
    if not raw_input:
        folder = resolve_path(ctx.task.workdir or ".")
        candidates = sorted(
            path
            for path in folder.iterdir()
            if path.suffix.lower() in (".xlsx", ".xlsm", ".xltx", ".xltm")
            and not path.name.startswith("~$")
            and not path.stem.endswith("_готово")
        )
        if not candidates:
            raise FileNotFoundError(f"В папке нет Excel-файла: {folder}")
        input_path = candidates[0]
    else:
        input_path = resolve_path(raw_input, ctx.task.workdir)
    if not input_path.exists():
        raise FileNotFoundError(f"Excel-файл не найден: {input_path}")
    output_raw = str(params.get("output_file", "")).strip()
    if output_raw:
        output_candidate = Path(output_raw)
        output_path = (
            output_candidate
            if output_candidate.is_absolute()
            else input_path.parent / output_candidate
        ).resolve()
    else:
        output_path = input_path.with_name(f"{input_path.stem}_готово.xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    ctx.log(f"Открытие: {input_path}")
    workbook = load_workbook(input_path)
    duration_name = str(params.get("duration_header", "Общая длительность"))
    mode = str(params.get("mode", "clean"))
    if mode == "format_only":
        changed = totals = 0
        for sheet in workbook.worksheets:
            ctx.check_cancelled()
            count, _valid = _normalize_sheet_durations(sheet)
            changed += count
            duration_column = _column(_header_map(sheet), duration_name)
            if duration_column:
                _add_total(sheet, duration_column)
                totals += 1
        workbook.save(output_path)
        message = f"Изменено значений времени: {changed}; листов с итогом: {totals}"
    else:
        output_book, sheet, max_column = _combine_sheets(workbook)
        headers = _header_map(sheet, max_column)
        name_name = str(params.get("name_header", "Название"))
        genre_name = str(params.get("genre_header", "Жанр"))
        name_column = _column(headers, name_name)
        if not name_column:
            raise ValueError(f"Не найден столбец «{name_name}»")
        genre_column = _column(headers, genre_name)
        removed = _clean_rows(
            sheet,
            max_column,
            name_column,
            genre_column,
            str(params.get("bad_value", "не указано")),
        )
        changed, _valid = _normalize_sheet_durations(sheet)
        duration_column = _column(_header_map(sheet, max_column), duration_name)
        total_text = "не рассчитан"
        if duration_column:
            _counted, total_text = _add_total(sheet, duration_column, max_column)
        output_book.save(output_path)
        message = (
            f"Удалено: «не указано» — {removed[0]}, неполных — {removed[1]}, "
            f"дублей — {removed[2]}; изменено времени — {changed}; итог — {total_text}"
        )
    ctx.log(f"Сохранено: {output_path}")
    return {"message": message, "files": [str(output_path)]}
